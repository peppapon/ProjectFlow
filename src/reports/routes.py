from flask import Blueprint, request, jsonify, send_file
from flask_jwt_extended import jwt_required, get_jwt_identity, get_jwt
from src.database.models import db, Task, User, Project, Report
from datetime import datetime
import os
import json

reports_bp = Blueprint('reports', __name__, url_prefix='/api/reports')

@reports_bp.route('/by-user/<int:user_id>', methods=['GET'])
@jwt_required()
def report_by_user(user_id):
    """Отчёт по исполнителю"""
    user = User.query.get(user_id)
    
    if not user:
        return jsonify({'error': 'User not found'}), 404
    
    tasks = Task.query.filter_by(assignee_id=user_id).all()
    
    # Группировка по статусам
    status_stats = {}
    for task in tasks:
        status_stats[task.status] = status_stats.get(task.status, 0) + 1
    
    return jsonify({
        'user': user.to_dict(),
        'total_tasks': len(tasks),
        'status_breakdown': status_stats,
        'tasks': [task.to_dict() for task in tasks]
    }), 200

@reports_bp.route('/by-project/<int:project_id>', methods=['GET'])
@jwt_required()
def report_by_project(project_id):
    """Отчёт по проекту"""
    project = Project.query.get(project_id)
    
    if not project:
        return jsonify({'error': 'Project not found'}), 404
    
    tasks = Task.query.filter_by(project_id=project_id).all()
    
    # Статистика по статусам
    status_stats = {}
    for task in tasks:
        status_stats[task.status] = status_stats.get(task.status, 0) + 1
    
    # Статистика по приоритетам
    priority_stats = {}
    for task in tasks:
        priority_stats[task.priority] = priority_stats.get(task.priority, 0) + 1
    
    return jsonify({
        'project': project.to_dict(),
        'total_tasks': len(tasks),
        'status_breakdown': status_stats,
        'priority_breakdown': priority_stats,
        'tasks': [task.to_dict() for task in tasks]
    }), 200

@reports_bp.route('/by-period', methods=['GET'])
@jwt_required()
def report_by_period():
    """Отчёт за период"""
    start_date = request.args.get('start_date')
    end_date = request.args.get('end_date')
    
    if not start_date or not end_date:
        return jsonify({'error': 'Start date and end date are required'}), 400
    
    try:
        start = datetime.fromisoformat(start_date.replace('Z', '+00:00'))
        end = datetime.fromisoformat(end_date.replace('Z', '+00:00'))
    except ValueError:
        return jsonify({'error': 'Invalid date format'}), 400
    
    # Задачи созданные в период
    tasks = Task.query.filter(
        Task.created_at >= start,
        Task.created_at <= end
    ).all()
    
    # Задачи завершенные в период
    completed_tasks = Task.query.filter(
        Task.completed_at >= start,
        Task.completed_at <= end
    ).all()
    
    return jsonify({
        'period': {
            'start': start_date,
            'end': end_date
        },
        'created_tasks': len(tasks),
        'completed_tasks': len(completed_tasks),
        'tasks': [task.to_dict() for task in tasks]
    }), 200

@reports_bp.route('/export/pdf', methods=['GET'])
@jwt_required()
def export_pdf():
    """Экспорт отчёта в PDF"""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    
    report_type = request.args.get('type', 'summary')
    filename = f'report_{datetime.now().strftime(\"%Y%m%d_%H%M%S\")}.pdf'
    filepath = os.path.join('reports', filename)
    
    # Создаем директорию если её нет
    os.makedirs('reports', exist_ok=True)
    
    # Создаем PDF документ
    doc = SimpleDocTemplate(filepath, pagesize=A4)
    elements = []
    
    styles = getSampleStyleSheet()
    
    # Заголовок
    title = Paragraph(f'ProjectFlow Report - {report_type.upper()}', styles['Title'])
    elements.append(title)
    elements.append(Spacer(1, 12))
    
    # Получаем данные для отчета
    tasks = Task.query.all()
    
    # Создаем таблицу
    data = [['ID', 'Title', 'Status', 'Priority', 'Assignee']]
    
    for task in tasks:
        assignee = User.query.get(task.assignee_id) if task.assignee_id else None
        data.append([
            str(task.id),
            task.title[:30],
            task.status,
            task.priority,
            assignee.username if assignee else 'Unassigned'
        ])
    
    table = Table(data)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
        ('GRID', (0, 0), (-1, -1), 1, colors.black)
    ]))
    
    elements.append(table)
    
    # Сохраняем отчет в БД
    current_user_id = get_jwt_identity()
    report = Report(
        name=f'PDF Report {datetime.now().strftime(\"%Y-%m-%d %H:%M\")}',
        report_type='pdf_export',
        parameters=json.dumps({'type': report_type}),
        generated_by=current_user_id,
        file_path=filepath
    )
    db.session.add(report)
    db.session.commit()
    
    # Генерируем PDF
    doc.build(elements)
    
    return send_file(filepath, as_attachment=True, download_name=filename)

@reports_bp.route('/export/excel', methods=['GET'])
@jwt_required()
def export_excel():
    """Экспорт отчёта в Excel"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    
    report_type = request.args.get('type', 'summary')
    filename = f'report_{datetime.now().strftime(\"%Y%m%d_%H%M%S\")}.xlsx'
    filepath = os.path.join('reports', filename)
    
    # Создаем директорию если её нет
    os.makedirs('reports', exist_ok=True)
    
    # Создаем Excel документ
    wb = Workbook()
    ws = wb.active
    ws.title = 'Tasks Report'
    
    # Заголовки
    headers = ['ID', 'Title', 'Status', 'Priority', 'Project', 'Assignee', 'Created At', 'Due Date']
    ws.append(headers)
    
    # Стилизация заголовков
    header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF')
    
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
    
    # Получаем данные
    tasks = Task.query.all()
    
    for task in tasks:
        assignee = User.query.get(task.assignee_id) if task.assignee_id else None
        project = Project.query.get(task.project_id)
        
        ws.append([
            task.id,
            task.title,
            task.status,
            task.priority,
            project.name if project else 'N/A',
            assignee.username if assignee else 'Unassigned',
            task.created_at.strftime('%Y-%m-%d %H:%M'),
            task.due_date.strftime('%Y-%m-%d') if task.due_date else 'N/A'
        ])
    
    # Автоматическая ширина колонок
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(cell.value)
            except:
                pass
        adjusted_width = (max_length + 2)
        ws.column_dimensions[column_letter].width = adjusted_width
    
    # Сохраняем
    wb.save(filepath)
    
    # Сохраняем отчет в БД
    current_user_id = get_jwt_identity()
    report = Report(
        name=f'Excel Report {datetime.now().strftime(\"%Y-%m-%d %H:%M\")}',
        report_type='excel_export',
        parameters=json.dumps({'type': report_type}),
        generated_by=current_user_id,
        file_path=filepath
    )
    db.session.add(report)
    db.session.commit()
    
    return send_file(filepath, as_attachment=True, download_name=filename)

@reports_bp.route('/summary', methods=['GET'])
@jwt_required()
def summary_report():
    """Сводный отчёт по всем задачам"""
    total_tasks = Task.query.count()
    total_projects = Project.query.count()
    total_users = User.query.count()
    
    # Задачи по статусам
    status_stats = {}
    for status in ['new', 'in_progress', 'review', 'completed', 'rejected']:
        count = Task.query.filter_by(status=status).count()
        status_stats[status] = count
    
    # Задачи по приоритетам
    priority_stats = {}
    for priority in ['low', 'medium', 'high']:
        count = Task.query.filter_by(priority=priority).count()
        priority_stats[priority] = count
    
    return jsonify({
        'summary': {
            'total_tasks': total_tasks,
            'total_projects': total_projects,
            'total_users': total_users
        },
        'status_breakdown': status_stats,
        'priority_breakdown': priority_stats
    }), 200
